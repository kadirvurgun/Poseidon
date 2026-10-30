import json,yaml,re,copy,sys,hashlib,datetime,subprocess
# Requirements: Python 3.10+, PyYAML, jsonschema (Draft 2020-12), openpyxl.
# Run from any directory: python 23.07_POSEIDON_Contract_Validation.py [package-root]
# This is a focused contract regression check, not a full OpenAPI validator.
if not __debug__: raise RuntimeError('Do not run validation with Python -O; assertions must be enabled.')
from pathlib import Path
from jsonschema import Draft202012Validator,FormatChecker
FORMAT_CHECKER=FormatChecker()
@FORMAT_CHECKER.checks('date-time', raises=(ValueError, TypeError))
def check_datetime(value):
 if not isinstance(value,str): return True
 if not re.fullmatch(r'\d{4}-\d{2}-\d{2}[Tt]\d{2}:\d{2}:\d{2}(?:\.\d+)?(?:[Zz]|[+-]\d{2}:\d{2})',value): return False
 return datetime.datetime.fromisoformat(value.upper().replace('Z','+00:00')).tzinfo is not None
R=Path(sys.argv[1]).resolve() if len(sys.argv)>1 else Path(__file__).resolve().parents[1]
checks=[]
def f(p):return next(R.rglob(p))
rule=json.loads(f('09.03*.json').read_text());Draft202012Validator.check_schema(rule);validator=Draft202012Validator(rule,format_checker=FORMAT_CHECKER);rules=json.loads(f('09.04*.json').read_text())
for r in rules:validator.validate(r)
negative=[{'all':[]},{'any':[]},{'field':'a','operator':'GT'},{'field':'a','operator':'GT','value':'x'},{'field':'a','operator':'IN','value':'x'},{'field':'a','operator':'DATE_LT','value':'2026-99-99'},{'field':'a','operator':'IS_TRUE','value':True},{'all':[{'field':'x','operator':'IS_TRUE'}],'surprise':True}]
for predicate in negative:
 r=copy.deepcopy(rules[0]);r['when']=predicate;assert list(validator.iter_errors(r)),predicate
for p in R.rglob('*.yaml'):
 d=yaml.safe_load(p.read_text());ids=[];ops=0
 def resolve(ref):
  assert ref.startswith('#/'),ref
  n=d
  for k in ref[2:].split('/'):n=n[k.replace('~1','/').replace('~0','~')]
  return n
 def walk(n):
  if isinstance(n,dict):
   if '$ref' in n:resolve(n['$ref'])
   for v in n.values():walk(v)
  elif isinstance(n,list):
   for v in n:walk(v)
 walk(d)
 for path,item in d['paths'].items():
  for method,op in item.items():
   ops+=1;ids.append(op['operationId']);assert d.get('security') or op.get('security');assert '401' in op['responses'] and '403' in op['responses']
   params=[resolve(x['$ref']) if '$ref' in x else x for x in op.get('parameters',[])];assert len({(x['name'],x['in']) for x in params})==len(params)
   for n in re.findall(r'{([^}]+)}',path):assert any(x['name']==n and x['in']=='path' and x['required'] for x in params)
   for code,resp in op['responses'].items():
    if int(code)>=400:assert 'application/problem+json' in resp['content']
 assert len(set(ids))==len(ids)
 for name,schema in d['components']['schemas'].items():Draft202012Validator.check_schema(schema)
 for sample,typ in [('09.05*.json','EvaluationRequest'),('09.06*.json','EvaluationResponse')]:
  schema={'$ref':'#/components/schemas/'+typ,'components':d['components']};Draft202012Validator(schema,format_checker=FORMAT_CHECKER).validate(json.loads(f(sample).read_text()))
 checks.append({'artifact':str(p.relative_to(R)),'operationCount':ops,'checks':['local_refs','unique_operation_ids','required_path_parameters','security_declaration','problem_errors','schema_structure','evaluation_samples'],'result':'PASS'})
 print(p.name,':',ops,'operations; refs, IDs, path parameters, security, errors, schemas and samples PASS')
# Check field identity against the controlled workbook, not only aggregate counts.
from openpyxl import load_workbook
workbook=load_workbook(f('14.02*.xlsx'),data_only=True,read_only=True)
rows=workbook['Data Dictionary 2316'].iter_rows(values_only=True);headers=next(rows);source_fields={}
for values in rows:
 row=dict(zip(headers,values));source_fields[(row['Screen ID'],row['Field Key'])]=row
product_for_fields=yaml.safe_load(f('16.01*.yaml').read_text())
field_checks=0
for name,schema in product_for_fields['components']['schemas'].items():
 for key,prop in schema.get('properties',{}).items():
  source=prop.get('x-source-field')
  if source:
   screen,field=source.split('.',1)
   assert key==field and (screen,field) in source_fields,(name,key,source)
   assert 'default' not in prop, 'Prototype defaults must not leak into domain DTOs'
   field_checks+=1
create_keys=set(product_for_fields['components']['schemas']['TugJobCreate']['properties'])
expected_create={field for (screen,field),row in source_fields.items() if screen=='tug-job-intake' and not str(row['Mandatory Rule'] or '').startswith('System')}
assert create_keys==expected_create
for name,screen in [('TugReadiness','tug-readiness'),('TugEquipmentReadiness','tug-equipment-readiness'),('TugCrewReadiness','tug-crew-readiness')]:
 assert set(product_for_fields['components']['schemas'][name]['properties'])=={field for sid,field in source_fields if sid==screen}
checks.append({'artifact':'14.02 to named API DTO field identity','checkedProperties':field_checks,'checks':['field_keys_match_source','create_excludes_system_fields','readiness_exact_source_key_sets','no_prototype_defaults'],'result':'PASS'})
# Typed contract regression: positive shape and malicious/malformed write rejection.
product=yaml.safe_load(f('16.01*.yaml').read_text())
def valid_named(name,payload):
 schema={'$ref':'#/components/schemas/'+name,'components':product['components']}
 return not list(Draft202012Validator(schema,format_checker=FORMAT_CHECKER).iter_errors(payload))
vessel={'company_id':'COMPANY-DEMO','name':'Demo vessel','flag':'TR','vessel_type':'ASD_TUG'}
job={'request_source':'INTERNAL','assisted_vessel':'Demo ship','job_type':'HARBOUR_ASSIST','requested_start':'2026-10-04T10:00:00Z','operation_area':'DEMO','required_tug_count':1,'priority':'NORMAL','request_owner':'USER-DEMO'}
certificate={'vessel_id':'VESSEL-DEMO','type':'DEMO_TYPE','issuer':'DEMO_ISSUER'}
verification={'evidence_id':'EVID-DEMO','result':'VERIFIED','reason':'Illustrative shape test only'}
assert valid_named('VesselWrite',vessel)
assert not valid_named('VesselWrite',{**vessel,'status':'READY'})
assert not valid_named('VesselWrite',{**vessel,'tenant_id':'OTHER-TENANT'})
assert not valid_named('VesselWrite',{**vessel,'imo_no':'123'})
assert valid_named('TugJobCreate',job)
assert not valid_named('TugJobCreate',{**job,'job_state':'READY'})
assert not valid_named('TugJobCreate',{**job,'required_tug_count':1.5})
assert not valid_named('TugJobCreate',{**job,'requested_start':'not-a-date'})
assert valid_named('CertificateWrite',certificate)
assert not valid_named('CertificateWrite',{**certificate,'verification_state':'VERIFIED'})
assert not valid_named('CertificateWrite',{**certificate,'expiry':'2026-99-99'})
assert valid_named('VerificationCommand',verification)
assert not valid_named('VerificationCommand',{'result':'VERIFIED','reason':'missing target'})
assert not valid_named('VerificationCommand',{**verification,'verifier':'FORGED'})
assert not valid_named('VerificationCommand',{**verification,'result':'READY'})
assert all(x.get('exampleOnly') is True and x.get('publicationStatus')=='DRAFT' and x.get('sourceVerificationStatus')=='PENDING' for x in rules)
for source_status,example in [('PENDING',False),('VERIFIED',True),('REJECTED',False)]:
 invalid=copy.deepcopy(rules[0]);invalid.update(publicationStatus='PUBLISHED',sourceVerificationStatus=source_status,exampleOnly=example)
 assert list(validator.iter_errors(invalid)), 'Unsafe publication fixture accepted'
# User-approved draft policy: incomplete persistence is separate from submission.
assert valid_named('VesselDraftSave',{'saveMode':'DRAFT','data':{}})
assert valid_named('TugJobDraftSave',{'saveMode':'DRAFT','data':{'assisted_vessel':'Demo'}})
assert valid_named('CertificateDraftSave',{'saveMode':'DRAFT','data':{}})
assert not valid_named('VesselWrite',{})
assert not valid_named('TugJobCreate',{'assisted_vessel':'Demo'})
assert not valid_named('CertificateWrite',{})
assert not valid_named('TugJobDraftSave',{'saveMode':'DRAFT','data':{'job_state':'READY'}})
assert not valid_named('CertificateDraftSave',{'saveMode':'DRAFT','data':{'verification_state':'VERIFIED'}})
assert not valid_named('TugJobDraftSave',{'saveMode':'DRAFT','data':{'required_tug_count':'Yes'}})
assert not valid_named('VesselDraftSave',{'saveMode':'SUBMIT','data':{}})
assert valid_named('DraftSaved',{'draftId':'DRAFT-DEMO','persistenceState':'DRAFT','submissionAllowed':False})
assert not valid_named('DraftSaved',{'draftId':'DRAFT-DEMO','persistenceState':'DRAFT','submissionAllowed':True})
for name in ['VesselDraftFields','TugJobDraftFields','CertificateDraftFields']:
 assert not any('default' in prop for prop in product['components']['schemas'][name]['properties'].values())
for path in ['/api/v1/vessels','/api/v1/vessels/{vesselId}','/api/v1/tug-jobs','/api/v1/certificates']:
 body=copy.deepcopy(product['paths'][path]['post']['requestBody']['content']['application/json']['schema']);body['components']=product['components']
 body_validator=Draft202012Validator(body,format_checker=FORMAT_CHECKER)
 assert not list(body_validator.iter_errors({'saveMode':'DRAFT','data':{}})),path
 assert list(body_validator.iter_errors({})),path
 response=copy.deepcopy(product['paths'][path]['post']['responses']['200']['content']['application/json']['schema']);response['components']=product['components']
 Draft202012Validator(response,format_checker=FORMAT_CHECKER).validate({'draftId':'DEMO','persistenceState':'DRAFT','submissionAllowed':False})
checks.append({'artifact':'User-approved default and draft policy','positiveDraftCases':4,'negativeDraftOrSubmissionCases':8,'checks':['incomplete_draft_allowed','incomplete_submission_rejected','no_yes_no_default_inference','draft_cannot_set_approval_or_ready','draft_ack_blocks_submission'],'result':'PASS','runtimeEnforcement':'NOT_VERIFIED'})
assert json.loads(f('09.06*.json').read_text())['simulation'] is True
checks.append({'artifact':'Named domain schemas and sample publication guard','positiveShapeCases':4,'negativeWriteCases':11,'negativePublicationCases':3,'checks':['server_state_injection_rejected','tenant_injection_rejected','date_and_count_types','human_verifier_spoof_rejected','draft_examples_not_publishable','simulation_label'],'result':'PASS','note':'Example controlled codes are illustrative; catalog lookup and authorization require server implementation.'})
r=json.loads(f('09.06*.json').read_text());assert {x['ruleId'] for x in r['results']}=={x['ruleId'] for x in rules}
for key,state in [('applicable','APPLICABLE'),('notApplicable','NOT_APPLICABLE'),('reviewRequired','REVIEW_REQUIRED')]:assert r['summary'][key]==sum(x['result']==state for x in r['results'])
print('6 rule samples, 8 malformed-predicate rejection tests, sample rule-set coverage and counts PASS')
print('STATIC CHECK ONLY: not deployed API behavior, legal approval, or full OpenAPI conformance certification.')

# Verify source/rule lineage and complete activation lists, not only schema validity.
by={x['ruleId']:x for x in rules}
for row in r['results']:
 source=by[row['ruleId']]
 assert row['ruleVersion']==source['version']
 assert row['sourceId']==source['source']['sourceId']
 assert row['sourceLocator']==source['source']['locator']
 assert row['activated']==([x['id'] for x in source['then'].get('activate',[])] if row['result']=='APPLICABLE' else [])
assert r['humanReview']['required'] is True and r['result']=='REVIEW_REQUIRED'
checks.append({'artifact':'Applicability DSL and samples','positiveRuleSamples':len(rules),'negativePredicateCases':len(negative),'checks':['rule_schema','reject_malformed_predicates','sample_rule_set_coverage','summary_counts','source_rule_version_lineage','full_activations','human_review_preserved'],'result':'PASS'})
gate=f('23.02*.sh')
subprocess.run(['bash','-n',str(gate)],check=True)
# No backend path: must fail safely before Docker/database mutation.
proc=subprocess.run(['bash',str(gate)],text=True,capture_output=True)
assert proc.returncode==2 and 'Usage:' in proc.stderr
checks.append({'artifact':str(gate.relative_to(R)),'checks':['bash_syntax','no_arguments_fail_closed_exit_2'],'result':'PASS'})
files=sorted(list(R.rglob('*.yaml'))+[p for p in R.rglob('*.json') if p.name.startswith('09.')]+[gate,Path(__file__).resolve(),f('14.02*.xlsx')])
report={'validationType':'FOCUSED_STATIC_CONTRACT_REGRESSION','status':'PASS','generatedAt':datetime.datetime.now(datetime.timezone.utc).isoformat(),'checks':checks,'limitations':['Not a full OpenAPI specification conformance validator.','No backend build, API server, Docker or database execution performed.','No authentication/authorization/idempotency/concurrency behavior proved.','Generic domain DTO refinement and critical state-transition contracts remain OPEN.','Rule structural validity does not establish legal applicability, source approval, or production readiness.'],'sha256':{(str(p.relative_to(R)) if p.is_relative_to(R) else 'validator/'+p.name):hashlib.sha256(p.read_bytes()).hexdigest() for p in files}}
out=R/'07_TEST_PILOT_RELEASE'/'23.07_POSEIDON_Contract_Validation.json';out.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n');print('Evidence:',out)
